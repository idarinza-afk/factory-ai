"""Engineering documents: machine list, part list, operation & maintenance manual.

Excel output uses openpyxl when installed (pip install improvise[docs]),
otherwise CSV. Prices that are not supplied are left blank on purpose -
the tool never invents prices.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from .tools import tool


def _out_dir() -> Path:
    from .config import load
    return load().output_dir


def _safe(name: str) -> str:
    return re.sub(r"[^\w\-]+", "_", name).strip("_") or "document"


def _write_table(path_stem: Path, title: str, headers: list[str], rows: list[list]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        path = path_stem.with_suffix(".csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([title])
            w.writerow(headers)
            w.writerows(rows)
        return path

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A3"
    path = path_stem.with_suffix(".xlsx")
    wb.save(path)
    return path


@tool
def make_machine_list(project: str, machines: list[dict]) -> dict:
    """Create a machine list (Excel/CSV) for a production line.

    Args:
        project: project or line name
        machines: list of {"name", "model" (opt), "qty", "capacity" (opt), "power_kw" (opt), "dimensions_mm" (opt), "supplier" (opt), "unit_price" (opt), "remark" (opt)}
    """
    headers = ["No.", "Machine", "Model", "Qty", "Capacity", "Power (kW)", "Dimensions L×W×H (mm)",
               "Supplier", "Unit price", "Total price", "Remark"]
    rows, total_kw = [], 0.0
    for i, m in enumerate(machines, start=1):
        qty = m.get("qty", 1)
        price = m.get("unit_price")
        total = price * qty if isinstance(price, (int, float)) else ""
        kw = m.get("power_kw", "")
        if isinstance(kw, (int, float)):
            total_kw += kw * qty
        rows.append([i, m.get("name", ""), m.get("model", ""), qty, m.get("capacity", ""), kw,
                     m.get("dimensions_mm", ""), m.get("supplier", ""), price if price is not None else "",
                     total, m.get("remark", "")])
    rows.append(["", "TOTAL CONNECTED LOAD", "", "", "", round(total_kw, 2), "", "", "", "", ""])
    path = _write_table(_out_dir() / f"{_safe(project)}_machine_list", f"Machine List - {project}", headers, rows)
    return {"file": str(path), "machines": len(machines), "total_power_kw": round(total_kw, 2)}


@tool
def make_part_list(machine: str, parts: list[dict]) -> dict:
    """Create a part list / spare-part list (Excel/CSV) for one machine.

    Args:
        machine: machine name
        parts: list of {"part_no", "description", "material" (opt), "qty", "unit" (opt), "maker" (opt), "spare_class" (opt: critical/wear/standard), "remark" (opt)}
    """
    headers = ["Item", "Part No.", "Description", "Material", "Qty", "Unit", "Maker / Brand", "Spare class", "Remark"]
    rows = [[i, p.get("part_no", ""), p.get("description", ""), p.get("material", ""), p.get("qty", 1),
             p.get("unit", "pcs"), p.get("maker", ""), p.get("spare_class", ""), p.get("remark", "")]
            for i, p in enumerate(parts, start=1)]
    path = _write_table(_out_dir() / f"{_safe(machine)}_part_list", f"Part List - {machine}", headers, rows)
    return {"file": str(path), "parts": len(parts)}


MANUAL_SECTIONS = [
    ("1. Safety", "safety"),
    ("2. Machine description", "description"),
    ("3. Technical specification", "specification"),
    ("4. Installation & utilities", "installation"),
    ("5. Operation", "operation"),
    ("6. Cleaning (CIP/COP) & hygiene", "cleaning"),
    ("7. Maintenance schedule", "maintenance"),
    ("8. Troubleshooting", "troubleshooting"),
    ("9. Spare parts", "spare_parts"),
]


@tool
def make_manual(machine: str, sections: dict) -> dict:
    """Assemble an operation & maintenance manual (Markdown, plus .docx if python-docx is installed).

    Args:
        machine: machine name
        sections: dict with any of: safety, description, specification, installation, operation, cleaning, maintenance, troubleshooting, spare_parts (text for each)
    """
    lines = [f"# Operation & Maintenance Manual - {machine}", ""]
    missing = []
    for title, key in MANUAL_SECTIONS:
        lines += [f"## {title}", ""]
        body = sections.get(key)
        if body:
            lines += [str(body), ""]
        else:
            lines += ["_To be completed._", ""]
            missing.append(key)
    md = "\n".join(lines)
    stem = _out_dir() / f"{_safe(machine)}_manual"
    md_path = stem.with_suffix(".md")
    md_path.write_text(md, encoding="utf-8")
    result = {"markdown": str(md_path), "missing_sections": missing}
    try:
        import docx  # python-docx
        d = docx.Document()
        for line in lines:
            if line.startswith("# "):
                d.add_heading(line[2:], 0)
            elif line.startswith("## "):
                d.add_heading(line[3:], 1)
            elif line:
                d.add_paragraph(line.strip("_"))
        d.save(stem.with_suffix(".docx"))
        result["docx"] = str(stem.with_suffix(".docx"))
    except ImportError:
        pass
    return result
